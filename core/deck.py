# -*- coding: utf-8 -*-
"""뷰 데이터 -> .pptx (python-pptx).

16:9 와이드. 모든 요소가 일반 도형 · 텍스트라서
PowerPoint 에서 자유롭게 다시 편집할 수 있습니다.
"""

from __future__ import annotations

import base64
import io
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

W = 13.333
H = 7.5

FONT = "맑은 고딕"

INK = RGBColor(0x14, 0x18, 0x1D)
INK2 = RGBColor(0x1D, 0x23, 0x29)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER2 = RGBColor(0xF4, 0xF6, 0xF8)
PAPER3 = RGBColor(0xE8, 0xEC, 0xF0)
LINE = RGBColor(0xD9, 0xDF, 0xE5)
TEXT2 = RGBColor(0x5A, 0x65, 0x6F)
TEXT3 = RGBColor(0x8B, 0x95, 0x9E)
INV2 = RGBColor(0xB3, 0xBE, 0xC8)


def _hex(s, default=(0x2C, 0x5F, 0x8A)):
    s = (s or "").lstrip("#")
    if len(s) == 6:
        try:
            return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))
        except ValueError:
            pass
    return RGBColor(*default)


# ---------------------------------------------------------------- 기본 도형
def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    sp.text_frame.text = ""
    return sp


def text(slide, x, y, w, h, blocks, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """blocks = [(문자열, {size, bold, color, space_after, line_spacing, font})]"""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor

    first = True
    for content, opt in blocks:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.line_spacing = opt.get("line_spacing", 1.35)
        p.space_after = Pt(opt.get("space_after", 0))
        for i, chunk in enumerate(str(content).split("\n")):
            if i:
                p.add_line_break()
            r = p.add_run()
            r.text = chunk
            r.font.size = Pt(opt.get("size", 12))
            r.font.bold = opt.get("bold", False)
            r.font.name = opt.get("font", FONT)
            r.font.color.rgb = opt.get("color", INK)
    return box


def picture(slide, data_uri, x, y, w, h):
    """data URI 이미지를 영역 안에 비율 유지해서 넣는다."""
    raw = _decode(data_uri)
    if not raw:
        return None
    try:
        from PIL import Image
        im = Image.open(io.BytesIO(raw))
        iw, ih = im.size
    except Exception:
        iw, ih = 4, 3
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        dw, dh = w, w / img_ratio
    else:
        dh, dw = h, h * img_ratio
    px = x + (w - dw) / 2
    py = y + (h - dh) / 2
    return slide.shapes.add_picture(io.BytesIO(raw), Inches(px), Inches(py), Inches(dw), Inches(dh))


def _decode(data_uri):
    if not data_uri:
        return None
    m = re.match(r"^data:[^;]+;base64,(.*)$", data_uri, re.S)
    if not m:
        return None
    try:
        return base64.b64decode(m.group(1))
    except Exception:
        return None


def est_lines(s, per_line):
    """대략적인 줄 수 (한글 기준). 레이아웃 높이를 정할 때 씁니다."""
    if not s:
        return 0
    total = 0
    for line in str(s).split("\n"):
        total += max(1, -(-len(line) // per_line))
    return total


# ---------------------------------------------------------------- 슬라이드
def cover(prs, v, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=INK)
    rect(s, 0, 0, 0.09, H, fill=accent)
    rect(s, W - 4.6, 0, 4.6, H, fill=INK2)

    p = v["person"]
    if p.get("photo"):
        picture(s, p["photo"], W - 4.0, 1.5, 3.4, 4.5)

    text(s, 0.85, 1.85, 7.6, 0.4,
         [(v["hero"]["eyebrow"].upper(), {"size": 11, "bold": True, "color": _lighten(accent)})])
    title = v["hero"]["title"] or p.get("name", "")
    tsize = 40 if len(title) < 34 else (33 if len(title) < 52 else 27)
    text(s, 0.85, 2.35, 7.6, 2.0,
         [(title, {"size": tsize, "bold": True, "color": WHITE, "line_spacing": 1.22})])
    if v["hero"]["tagline"]:
        text(s, 0.85, 4.55, 7.4, 1.4,
             [(v["hero"]["tagline"], {"size": 12, "color": INV2, "line_spacing": 1.5})])

    rect(s, 0.85, 6.25, 7.4, 0.02, fill=RGBColor(0x38, 0x41, 0x4B))
    foot = " · ".join(x for x in [p.get("name", ""), p.get("email", ""), p.get("phone", "")] if x)
    text(s, 0.85, 6.45, 7.4, 0.4, [(foot, {"size": 11, "color": INV2})])
    return s


def _lighten(c):
    return RGBColor(min(255, c[0] + 70), min(255, c[1] + 70), min(255, c[2] + 60))


def _header(s, label, title, accent):
    rect(s, 0.85, 0.5, 0.32, 0.045, fill=accent)
    text(s, 0.85, 0.62, 6.0, 0.32, [(label.upper(), {"size": 10.5, "bold": True, "color": accent})])
    text(s, 0.85, 0.95, 9.0, 0.5, [(title, {"size": 22, "bold": True, "color": INK})])
    rect(s, 0, 1.62, W, 0.012, fill=LINE)


def profile(prs, v, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=WHITE)
    _header(s, "Profile", "기본 정보", accent)

    p = v["person"]
    if p.get("photo"):
        picture(s, p["photo"], 0.85, 2.0, 2.5, 3.3)
        rect(s, 0.85, 5.45, 2.5, 0.5, fill=INK)
        text(s, 0.95, 5.58, 2.3, 0.3,
             [("%s  %s" % (p.get("name", ""), p.get("name_en", "")),
               {"size": 12, "bold": True, "color": WHITE})])
        left = 3.9
    else:
        left = 0.85

    y = 2.0
    rows = [r for r in [
        ("생년월일", p.get("birth", "")),
        ("연락처", p.get("phone", "")),
        ("이메일", p.get("email", "")),
        ("거주지", p.get("location", "")),
    ] if r[1]]
    if rows:
        text(s, left, y, 3.0, 0.3, [("CONTACT", {"size": 10, "bold": True, "color": TEXT3})])
        y += 0.38
        for k, val in rows:
            text(s, left, y, 1.1, 0.3, [(k, {"size": 10.5, "color": TEXT3})])
            text(s, left + 1.15, y, 3.4, 0.3, [(val, {"size": 11.5, "color": INK})])
            y += 0.33
        y += 0.25

    if v["education"]:
        text(s, left, y, 3.0, 0.3, [("EDUCATION", {"size": 10, "bold": True, "color": TEXT3})])
        y += 0.38
        for ed in v["education"][:3]:
            text(s, left, y, 4.6, 0.3,
                 [(ed.get("school", ""), {"size": 11.5, "bold": True, "color": INK})])
            sub = " · ".join(x for x in [ed.get("dept", ""), ed.get("period", "")] if x)
            text(s, left, y + 0.24, 4.6, 0.3, [(sub, {"size": 10, "color": TEXT2})])
            y += 0.62

    # 오른쪽: 경력
    rx = left + 5.0 if left > 1 else 6.4
    ry = 2.0
    if v["experience"]:
        text(s, rx, ry, 3.0, 0.3, [("EXPERIENCE", {"size": 10, "bold": True, "color": TEXT3})])
        ry += 0.4
        for ex in v["experience"][:4]:
            rect(s, rx, ry, 0.035, 0.62, fill=accent)
            org = re.sub(r"\s*\(.*?\)", "", ex.get("org", ""))
            text(s, rx + 0.2, ry, 5.6, 0.3, [(org, {"size": 12, "bold": True, "color": INK})])
            line2 = " · ".join(x for x in [ex.get("role", ""), ex.get("period", "")] if x)
            text(s, rx + 0.2, ry + 0.26, 5.6, 0.3, [(line2, {"size": 10, "color": TEXT2})])
            ry += 0.8

    if v["languages"]:
        text(s, rx, ry + 0.1, 3.0, 0.3, [("LANGUAGES", {"size": 10, "bold": True, "color": TEXT3})])
        ry += 0.5
        for lg in v["languages"][:4]:
            label = lg.get("name", "")
            detail = " ".join(x for x in [lg.get("level", ""), lg.get("cert", "")] if x)
            text(s, rx, ry, 5.6, 0.3, [("%s   %s" % (label, detail), {"size": 11, "color": INK})])
            ry += 0.3
    return s


def index_slide(prs, v, plan, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=PAPER2)
    _header(s, "Index", "목차", accent)
    if v["index_note"]:
        text(s, 0.85, 1.9, 8.0, 0.3, [(v["index_note"], {"size": 11, "color": TEXT2})])

    top = 2.45
    bottom = 6.8
    n = max(1, len(plan))
    step = min(1.0, (bottom - top) / n)
    dense = step < 0.62
    for i, (_sid, label, desc) in enumerate(plan):
        y = top + i * step
        text(s, 0.85, y, 0.7, 0.3,
             [("%02d" % (i + 1), {"size": 13 if not dense else 11, "bold": True, "color": accent})])
        text(s, 1.6, y - 0.02, 5.4, 0.3,
             [(label, {"size": 15 if not dense else 12.5, "bold": True, "color": INK})])
        text(s, 1.6, y + 0.26 if not dense else y + 0.22, 8.0, 0.3,
             [(desc, {"size": 10.5 if not dense else 9.5, "color": TEXT2})])
        rect(s, 0.85, y + step - 0.12, W - 1.7, 0.008, fill=LINE)
    return s


def project_slide(prs, pr, num, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=WHITE)

    # 상단
    rect(s, 0, 0, W, 1.75, fill=PAPER2)
    rect(s, 0, 1.74, W, 0.012, fill=LINE)
    text(s, 0.85, 0.42, 3.0, 0.3,
         [("PROJECT %02d" % num, {"size": 10.5, "bold": True, "color": accent})])
    title = pr.get("title", "")
    text(s, 0.85, 0.72, 8.4, 0.5,
         [(title, {"size": 23 if len(title) < 30 else 19, "bold": True, "color": INK})])
    sub = pr.get("subtitle", "")
    if sub:
        text(s, 0.85, 1.18, 8.4, 0.35, [(sub, {"size": 12, "bold": True, "color": accent})])
    meta = " · ".join(x for x in [pr.get("org", ""), pr.get("period", ""), pr.get("my_role", "")] if x)
    if meta:
        text(s, 0.85, 1.44, 8.4, 0.3, [(meta, {"size": 9.5, "color": TEXT3})])

    if pr.get("metric_value"):
        text(s, W - 3.6, 0.72, 2.75, 0.55,
             [(pr["metric_value"], {"size": 27, "bold": True, "color": INK})], align=PP_ALIGN.RIGHT)
        text(s, W - 3.6, 1.28, 2.75, 0.3,
             [(pr.get("metric_label", ""), {"size": 10, "color": TEXT3})], align=PP_ALIGN.RIGHT)

    # 하단 성과 박스를 먼저 잡아야 본문이 그 위로 넘치지 않는다
    result = pr.get("result", "")
    if result:
        rh = max(0.85, 0.24 * est_lines(result, 88) + 0.5)
        result_top = H - rh - 0.35
    else:
        rh, result_top = 0, H - 0.35
    floor = result_top - 0.15          # 본문이 넘어가면 안 되는 선

    # 본문 좌: 배경 + 실행
    y = 2.1
    bg = pr.get("background", "")
    if bg:
        text(s, 0.85, y, 1.5, 0.25, [("BACKGROUND", {"size": 9.5, "bold": True, "color": TEXT3})])
        bh = 0.24 * est_lines(bg, 42) + 0.1
        text(s, 0.85, y + 0.3, 7.0, bh,
             [(bg, {"size": 10.5, "color": TEXT2, "line_spacing": 1.5})])
        y += 0.3 + bh + 0.24

    actions = [a for a in (pr.get("actions") or []) if a]
    if actions:
        text(s, 0.85, y, 1.5, 0.25, [("ACTION", {"size": 9.5, "bold": True, "color": TEXT3})])
        y += 0.32
        many = len(actions) >= 6
        asize = 9.5 if many else 10.5
        astep = 0.30 if many else 0.34
        wrap = 52 if many else 46
        extras = [0.22 * max(0, est_lines(a, wrap) - 1) for a in actions]

        # 남은 높이에 안 들어가면 줄 간격과 글자를 함께 줄인다
        need = sum(astep + e for e in extras)
        room = floor - y
        if need > room > 0:
            shrink = max(0.72, room / need)
            astep *= shrink
            extras = [e * shrink for e in extras]
            asize = max(8.0, asize * shrink)

        for i, a in enumerate(actions):
            text(s, 0.85, y, 0.35, 0.25,
                 [("%02d" % (i + 1), {"size": 8.5, "bold": True, "color": accent})])
            text(s, 1.25, y - 0.02, 6.6, 0.3 + extras[i],
                 [(a, {"size": asize, "color": INK, "line_spacing": 1.4})])
            y += astep + extras[i]

    # 본문 우: 이미지 + 핵심 역량
    skills = pr.get("core_skills") or []
    img_bottom = 4.65
    if skills:
        # CORE 줄이 들어갈 자리를 먼저 떼어 놓고 남는 만큼만 이미지에 준다
        want = 0.3 + 0.35 * min(len(skills), 5)
        img_bottom = min(4.65, floor - want)
    if pr.get("image") and img_bottom > 2.6:
        picture(s, pr["image"], 8.35, 2.15, 4.1, img_bottom - 2.15)

    ry = (img_bottom + 0.2) if pr.get("image") else 2.2
    if skills:
        text(s, 8.35, ry, 2.0, 0.25, [("CORE", {"size": 9.5, "bold": True, "color": TEXT3})])
        ry += 0.3
        for sk in skills[:5]:
            if ry + 0.3 > floor:       # 성과 박스를 침범하면 멈춘다
                break
            rect(s, 8.35, ry, 4.1, 0.3, fill=PAPER3)
            text(s, 8.5, ry + 0.05, 3.9, 0.25, [(sk, {"size": 9.5, "color": INK})])
            ry += 0.35

    if result:
        rect(s, 0.85, result_top, W - 1.7, rh, fill=INK)
        text(s, 1.1, result_top + 0.22, 1.0, 0.25,
             [("RESULT", {"size": 9.5, "bold": True, "color": _lighten(accent)})])
        text(s, 2.2, result_top + 0.2, W - 3.4, rh - 0.35,
             [(result, {"size": 11, "color": WHITE, "line_spacing": 1.5})])
    return s


def custom_slide(prs, custom, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=INK)
    rect(s, 0.85, 0.62, 0.32, 0.045, fill=accent)
    text(s, 0.85, 0.74, 6.0, 0.32,
         [("PERSPECTIVE", {"size": 10.5, "bold": True, "color": _lighten(accent)})])
    text(s, 0.85, 1.06, 9.0, 0.5,
         [(custom.get("title") or "직무에 대한 관점", {"size": 22, "bold": True, "color": WHITE})])
    if custom.get("subtitle"):
        text(s, 0.85, 1.56, 10.5, 0.3,
             [(custom["subtitle"], {"size": 10.5, "color": INV2})])

    items = custom["items"][:3]
    y = 2.15
    h = (H - 2.65) / max(1, len(items))
    for it in items:
        rect(s, 0.85, y, W - 1.7, h - 0.22, fill=INK2)
        rect(s, 0.85, y, 0.04, h - 0.22, fill=accent)
        text(s, 1.15, y + 0.22, W - 2.4, 0.3,
             [(it.get("q", ""), {"size": 12.5, "bold": True, "color": WHITE})])
        text(s, 1.15, y + 0.6, W - 2.4, h - 0.9,
             [(it.get("a", ""), {"size": 10.5, "color": INV2, "line_spacing": 1.5})])
        y += h
    return s


def skills_slide(prs, v, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=PAPER2)
    _header(s, "Skills & Languages", "보유 역량", accent)

    def bars(items, x, y, width, label_of, detail_of):
        for it in items:
            text(s, x, y, width * 0.6, 0.25,
                 [(label_of(it), {"size": 11, "color": INK})])
            d = detail_of(it)
            if d:
                text(s, x, y, width, 0.25, [(d, {"size": 9, "color": TEXT3})], align=PP_ALIGN.RIGHT)
            rect(s, x, y + 0.27, width, 0.075, fill=PAPER3)
            pct = max(0, min(100, int(it.get("bar") or 80)))
            if pct:
                rect(s, x, y + 0.27, width * pct / 100.0, 0.075, fill=accent)
            y += 0.52
        return y

    y = 2.05
    if v["skills"]:
        text(s, 0.85, y, 3.0, 0.25, [("TOOLS & SKILLS", {"size": 10, "bold": True, "color": TEXT3})])
        bars(v["skills"][:8], 0.85, y + 0.35, 5.1, lambda i: i.get("name", ""), lambda i: "")

    y = 2.05
    if v["languages"]:
        text(s, 7.0, y, 3.0, 0.25, [("LANGUAGES", {"size": 10, "bold": True, "color": TEXT3})])
        y = bars(v["languages"][:4], 7.0, y + 0.35, 5.4,
                 lambda i: i.get("name", ""),
                 lambda i: " ".join(x for x in [i.get("level", ""), i.get("cert", "")] if x))
    if v["certificates"]:
        text(s, 7.0, y + 0.15, 3.0, 0.25, [("CERTIFICATES", {"size": 10, "bold": True, "color": TEXT3})])
        y += 0.5
        for c in v["certificates"][:4]:
            rect(s, 7.0, y, 5.4, 0.42, fill=WHITE, line=LINE)
            text(s, 7.15, y + 0.09, 2.3, 0.25, [(c.get("name", ""), {"size": 10.5, "color": INK})])
            sub = " · ".join(x for x in [c.get("org", ""), c.get("date", "")] if x)
            text(s, 9.5, y + 0.11, 2.75, 0.25, [(sub, {"size": 8.5, "color": TEXT3})],
                 align=PP_ALIGN.RIGHT)
            y += 0.5
    return s


def more_slide(prs, v, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=WHITE)
    _header(s, "Education & Activities", "학력 · 수상 · 활동", accent)

    cols = []
    if v["education"]:
        cols.append(("EDUCATION", [(x.get("school", ""),
                                    " · ".join(y for y in [x.get("dept", ""), x.get("period", "")] if y))
                                   for x in v["education"]]))
    if v["awards"]:
        cols.append(("AWARDS", [(x.get("name", ""),
                                 " · ".join(y for y in [x.get("prize", ""), x.get("date", "")] if y))
                                for x in v["awards"]]))
    if v["activities"]:
        cols.append(("ACTIVITIES", [(x.get("name", ""),
                                     " · ".join(y for y in [x.get("type", ""), x.get("period", "")] if y))
                                    for x in v["activities"]]))
    if not cols:
        return s

    width = (W - 1.7 - 0.6 * (len(cols) - 1)) / len(cols)
    for i, (head, items) in enumerate(cols):
        x = 0.85 + i * (width + 0.6)
        text(s, x, 2.0, width, 0.3, [(head, {"size": 10, "bold": True, "color": accent})])
        rect(s, x, 2.32, width, 0.02, fill=INK)
        y = 2.5
        for name, sub in items[:6]:
            text(s, x, y, width, 0.3, [(name, {"size": 11.5, "bold": True, "color": INK})])
            if sub:
                text(s, x, y + 0.26, width, 0.3, [(sub, {"size": 9.5, "color": TEXT2})])
            y += 0.66
    return s


def closing(prs, v, accent):
    s = _blank(prs)
    rect(s, 0, 0, W, H, fill=INK)
    rect(s, W / 2 - 0.2, 2.5, 0.4, 0.045, fill=accent)
    text(s, 1.0, 2.9, W - 2.0, 0.8,
         [("읽어주셔서 감사합니다.", {"size": 30, "bold": True, "color": WHITE})], align=PP_ALIGN.CENTER)
    p = v["person"]
    foot = "   ".join(x for x in [p.get("name", ""), p.get("email", ""), p.get("phone", "")] if x)
    text(s, 1.0, 4.0, W - 2.0, 0.4, [(foot, {"size": 12, "color": INV2})], align=PP_ALIGN.CENTER)
    return s


# ---------------------------------------------------------------- 전체
def build(v):
    """뷰 데이터로 프레젠테이션을 만들어 BytesIO 로 돌려준다."""
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    accent = _hex(v.get("accent"))

    plan = []
    if v["experience"]:
        plan.append(("experience", "Experience", "경력 사항"))
    if v["projects"]:
        plan.append(("projects", "Projects", "주요 프로젝트"))
    if v["custom"]:
        plan.append(("custom", v["custom"].get("title") or "Perspective",
                     v["custom"].get("subtitle") or "직무에 대한 관점"))
    if v["skills"] or v["languages"]:
        plan.append(("skills", "Skills & Languages", "보유 역량"))
    if v["education"] or v["awards"] or v["activities"]:
        plan.append(("more", "Education & Activities", "학력 · 수상 · 활동"))

    cover(prs, v, accent)
    profile(prs, v, accent)
    if plan:
        index_slide(prs, v, plan, accent)
    for i, pr in enumerate(v["projects"]):
        project_slide(prs, pr, i + 1, accent)
    if v["custom"]:
        custom_slide(prs, v["custom"], accent)
    if v["skills"] or v["languages"] or v["certificates"]:
        skills_slide(prs, v, accent)
    if v["education"] or v["awards"] or v["activities"]:
        more_slide(prs, v, accent)
    closing(prs, v, accent)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf
